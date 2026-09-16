import fitz  # PyMuPDF 
import docx
import io


def extract_text_from_file(file_content: bytes, filename: str) -> str:
    file_ext = filename.lower().split('.')[-1]

    if file_ext == "pdf":
        return _extract_from_pdf(file_content)
    elif file_ext in ["doc", "docx"]:
        return _extract_from_docx(file_content)
    elif file_ext == "txt":
        return " ".join(file_content.decode("utf-8", errors="ignore").split())
    else:
        raise ValueError(f"Unsupported file type: .{file_ext}. Upload PDF, DOCX, or TXT.")


def extract_text_from_plain(text: str) -> str:
    # for when recruiter pastes JD directly as text
    return " ".join(text.split())


def _extract_from_pdf(file_content: bytes) -> str:
    text = ""
    with fitz.open(stream=file_content, filetype="pdf") as doc:
        for page in doc:
            text += page.get_text() + "\n"
    return " ".join(text.split())


def _extract_from_pdf_raw(file_content: bytes) -> str:
    """Return PDF text with newlines preserved (used for name extraction)."""
    text = ""
    with fitz.open(stream=file_content, filetype="pdf") as doc:
        for page in doc:
            text += page.get_text() + "\n"
    return text


def _extract_from_docx(file_content: bytes) -> str:
    doc = docx.Document(io.BytesIO(file_content))
    parts = []
    # body paragraphs
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)
    # text inside tables (many DOCX JDs/resumes use table-based layouts)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t:
                    parts.append(t)
    return " ".join(parts)


def extract_raw_text_for_name(file_content: bytes, filename: str) -> str:
    """Extract text preserving line breaks — only used for candidate name extraction."""
    file_ext = filename.lower().split('.')[-1]
    if file_ext == "pdf":
        return _extract_from_pdf_raw(file_content)
    elif file_ext in ["doc", "docx"]:
        doc = docx.Document(io.BytesIO(file_content))
        parts = [p.text for p in doc.paragraphs]
        # also grab table cells so name detection works for table-based layouts
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)
    elif file_ext == "txt":
        return file_content.decode("utf-8", errors="ignore")
    return ""
